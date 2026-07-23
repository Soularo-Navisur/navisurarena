# routes/finances.py — NAVISUR v9.9
# Ce fichier est un module de routes NAVISUR.
# Il importe tout depuis core.py (app, get_db, nlog, etc.)
from core import *

@app.route('/commissions')
def commissions_list():
    conn = get_db()
    annee = request.args.get('annee', datetime.now().year)
    mois = request.args.get('mois', '')
    statut = request.args.get('statut', '')
    query = """SELECT cm.*, ct.numero, ct.type_assurance, c.nom, c.prenom
               FROM commissions cm LEFT JOIN contrats ct ON cm.contrat_id = ct.id
               LEFT JOIN clients c ON ct.client_id = c.id WHERE 1=1"""
    params = []
    if annee:
        query += " AND cm.annee = ?"
        params.append(int(annee))
    if mois:
        query += " AND cm.mois = ?"
        params.append(int(mois))
    if statut:
        query += " AND cm.statut = ?"
        params.append(statut)
    query += " ORDER BY cm.annee DESC, cm.mois DESC"
    commissions = conn.execute(query, params).fetchall()
    total = conn.execute("SELECT COALESCE(SUM(montant_net),0) FROM commissions WHERE annee=?",
                         (int(annee),)).fetchone()[0]
    total_recu = conn.execute(
        "SELECT COALESCE(SUM(montant_net),0) FROM commissions WHERE annee=? AND statut='recu'",
        (int(annee),)).fetchone()[0]
    conn.close()
    return render_template('commissions/index.html', commissions=commissions,
                           annee=annee, mois=mois, statut=statut, total=total,
                           total_recu=total_recu,
                           annees=list(range(2020, datetime.now().year + 2)))

@app.route('/commissions/nouvelle', methods=['GET', 'POST'])
def commission_nouvelle():
    conn = get_db()
    if request.method == 'POST':
        montant_brut = float(request.form.get('montant_brut') or 0)
        taux = float(request.form.get('taux') or 0)
        montant_net = montant_brut * (taux / 100) if taux > 0 else montant_brut
        conn.execute('''INSERT INTO commissions (contrat_id, compagnie, mois, annee,
            montant_brut, taux, montant_net, statut, date_paiement, notes)
            VALUES (?,?,?,?,?,?,?,?,?,?)''',
            (request.form.get('contrat_id') or None, get_compagnie_from_form(),
             int(request.form.get('mois', datetime.now().month)),
             int(request.form.get('annee', datetime.now().year)),
             montant_brut, taux, montant_net,
             request.form.get('statut', 'attendu'), request.form.get('date_paiement', ''),
             request.form.get('notes', '')))
        conn.commit()
        conn.close()
        flash('Commission enregistrée !', 'success')
        return redirect(url_for('commissions_list'))
    contrats = conn.execute("""SELECT ct.id, ct.numero, ct.compagnie, c.nom
                               FROM contrats ct JOIN clients c ON ct.client_id=c.id
                               ORDER BY c.nom""").fetchall()
    conn.close()
    return render_template('commissions/form.html', commission=None, contrats=contrats,
                           titre='Nouvelle commission', now=datetime.now())

@app.route('/commissions/<int:id>/modifier', methods=['GET', 'POST'])
def commission_modifier(id):
    conn = get_db()
    commission = conn.execute("SELECT * FROM commissions WHERE id=?", (id,)).fetchone()
    if request.method == 'POST':
        montant_brut = float(request.form.get('montant_brut') or 0)
        taux = float(request.form.get('taux') or 0)
        montant_net = montant_brut * (taux / 100) if taux > 0 else float(
            request.form.get('montant_net') or montant_brut)
        conn.execute('''UPDATE commissions SET contrat_id=?, compagnie=?, mois=?, annee=?,
            montant_brut=?, taux=?, montant_net=?, statut=?, date_paiement=?, notes=? WHERE id=?''',
            (request.form.get('contrat_id') or None, get_compagnie_from_form(),
             int(request.form.get('mois', 1)), int(request.form.get('annee', date.today().year)),
             montant_brut, taux, montant_net,
             request.form.get('statut', 'attendu'), request.form.get('date_paiement', ''),
             request.form.get('notes', ''), id))
        conn.commit()
        conn.close()
        flash('Commission mise à jour !', 'success')
        return redirect(url_for('commissions_list'))
    contrats = conn.execute("""SELECT ct.id, ct.numero, ct.compagnie, c.nom
                               FROM contrats ct JOIN clients c ON ct.client_id=c.id
                               ORDER BY c.nom""").fetchall()
    conn.close()
    return render_template('commissions/form.html', commission=commission, contrats=contrats,
                           titre='Modifier la commission', now=datetime.now())

@app.route('/commissions/<int:id>/supprimer', methods=['POST'])
def commission_supprimer(id):
    conn = get_db()
    conn.execute("DELETE FROM commissions WHERE id=?", (id,))
    conn.commit()
    conn.close()
    flash('Commission supprimée.', 'info')
    return redirect(url_for('commissions_list'))

# ─── FACTURES ────────────────────────────────────────────────────────────────

@app.route('/factures')
def factures_list():
    conn = get_db()
    statut = request.args.get('statut', '')
    query = """SELECT f.*, c.nom, c.prenom FROM factures f
               JOIN clients c ON f.client_id = c.id WHERE 1=1"""
    params = []
    if statut:
        query += " AND f.statut = ?"
        params.append(statut)
    query += " ORDER BY f.date_emission DESC"
    factures = conn.execute(query, params).fetchall()
    total_ttc = conn.execute(
        "SELECT COALESCE(SUM(montant_ttc),0) FROM factures WHERE statut='envoyee'").fetchone()[0]
    conn.close()
    return render_template('factures/index.html', factures=factures, statut=statut,
                           total_impaye=total_ttc)

@app.route('/factures/nouvelle', methods=['GET', 'POST'])
def facture_nouvelle():
    conn = get_db()
    if request.method == 'POST':
        montant_ht = float(request.form.get('montant_ht') or 0)
        tva = float(request.form.get('tva') or 20)
        montant_ttc = montant_ht * (1 + tva / 100)
        annee = datetime.now().year
        count = conn.execute("SELECT COUNT(*) FROM factures").fetchone()[0] + 1
        numero = f"FAC-{annee}-{count:04d}"
        conn.execute('''INSERT INTO factures (client_id, numero, objet, date_emission,
            date_echeance, montant_ht, tva, montant_ttc, statut, notes, iban)
            VALUES (?,?,?,?,?,?,?,?,?,?,?)''',
            (request.form['client_id'], numero, request.form.get('objet', ''),
             request.form.get('date_emission', date.today().isoformat()),
             request.form.get('date_echeance', ''),
             montant_ht, tva, montant_ttc,
             request.form.get('statut', 'brouillon'), request.form.get('notes', ''),
             request.form.get('iban', '')))
        conn.commit()
        conn.close()
        flash(f'Quittance {numero} créée !', 'success')
        return redirect(url_for('factures_list'))
    clients = conn.execute("SELECT id, nom, prenom FROM clients ORDER BY nom").fetchall()
    conn.close()
    return render_template('factures/form.html', facture=None, clients=clients,
                           titre='Nouvelle quittance', today=date.today().isoformat())

@app.route('/factures/<int:id>/modifier', methods=['GET', 'POST'])
def facture_modifier(id):
    conn = get_db()
    facture = conn.execute("SELECT * FROM factures WHERE id=?", (id,)).fetchone()
    if request.method == 'POST':
        montant_ht = float(request.form.get('montant_ht') or 0)
        tva = float(request.form.get('tva') or 20)
        montant_ttc = montant_ht * (1 + tva / 100)
        conn.execute('''UPDATE factures SET client_id=?, objet=?, date_emission=?, date_echeance=?,
            montant_ht=?, tva=?, montant_ttc=?, statut=?, notes=?, iban=? WHERE id=?''',
            (request.form['client_id'], request.form.get('objet', ''),
             request.form.get('date_emission', ''), request.form.get('date_echeance', ''),
             montant_ht, tva, montant_ttc,
             request.form.get('statut', 'brouillon'), request.form.get('notes', ''),
             request.form.get('iban', ''), id))
        conn.commit()
        conn.close()
        flash('Facture mise à jour !', 'success')
        return redirect(url_for('factures_list'))
    clients = conn.execute("SELECT id, nom, prenom FROM clients ORDER BY nom").fetchall()
    conn.close()
    return render_template('factures/form.html', facture=facture, clients=clients,
                           titre='Modifier la facture', today=date.today().isoformat())

@app.route('/factures/<int:id>/supprimer', methods=['POST'])
def facture_supprimer(id):
    conn = get_db()
    conn.execute("DELETE FROM factures WHERE id=?", (id,))
    conn.commit()
    conn.close()
    flash('Quittance supprimée.', 'info')
    return redirect(url_for('factures_list'))

@app.route('/factures/<int:id>/imprimer')
def facture_imprimer(id):
    conn = get_db()
    facture = conn.execute("""SELECT f.*, c.nom, c.prenom, c.adresse, c.code_postal,
                              c.ville, c.email, c.telephone, c.siren
                              FROM factures f JOIN clients c ON f.client_id=c.id
                              WHERE f.id=?""", (id,)).fetchone()
    conn.close()
    if not facture:
        flash('Quittance introuvable.', 'error')
        return redirect(url_for('factures_list'))
    return render_template('factures/print.html', facture=facture, now=datetime.now())

# ─── TÂCHES ──────────────────────────────────────────────────────────────────

@app.route('/taches')
def taches_list():
    conn = get_db()
    statut = request.args.get('statut', 'a_faire')
    priorite = request.args.get('priorite', '')
    query = """SELECT t.*, c.nom, c.prenom FROM taches t
               LEFT JOIN clients c ON t.client_id = c.id WHERE 1=1"""
    params = []
    if statut:
        query += " AND t.statut = ?"
        params.append(statut)
    if priorite:
        query += " AND t.priorite = ?"
        params.append(priorite)
    query += """ ORDER BY CASE t.priorite WHEN 'urgente' THEN 1 WHEN 'haute' THEN 2
                WHEN 'normale' THEN 3 ELSE 4 END, t.date_echeance ASC"""
    taches = conn.execute(query, params).fetchall()
    counts = {
        'a_faire': conn.execute("SELECT COUNT(*) FROM taches WHERE statut='a_faire'").fetchone()[0],
        'en_cours': conn.execute("SELECT COUNT(*) FROM taches WHERE statut='en_cours'").fetchone()[0],
        'terminee': conn.execute("SELECT COUNT(*) FROM taches WHERE statut='terminee'").fetchone()[0],
        'en_retard': conn.execute(
            "SELECT COUNT(*) FROM taches WHERE statut='a_faire' AND date_echeance < ?",
            (date.today().isoformat(),)).fetchone()[0],
    }
    conn.close()
    return render_template('taches/index.html', taches=taches, statut=statut,
                           priorite=priorite, counts=counts, today=date.today().isoformat())

@app.route('/taches/nouvelle', methods=['GET', 'POST'])
def tache_nouvelle():
    conn = get_db()
    if request.method == 'POST':
        conn.execute('''INSERT INTO taches (titre, description, client_id, contrat_id,
            date_echeance, priorite, statut) VALUES (?,?,?,?,?,?,?)''',
            (request.form['titre'], request.form.get('description', ''),
             request.form.get('client_id') or None, request.form.get('contrat_id') or None,
             request.form.get('date_echeance', ''), request.form.get('priorite', 'normale'),
             request.form.get('statut', 'a_faire')))
        conn.commit()
        conn.close()
        flash('Tâche créée !', 'success')
        return redirect(url_for('taches_list'))
    clients = conn.execute("SELECT id, nom, prenom FROM clients ORDER BY nom").fetchall()
    contrats = conn.execute("""SELECT ct.id, ct.numero, ct.type_assurance, c.nom
                               FROM contrats ct JOIN clients c ON ct.client_id=c.id
                               ORDER BY c.nom""").fetchall()
    conn.close()
    return render_template('taches/form.html', tache=None, clients=clients, contrats=contrats,
                           titre='Nouvelle tâche', preselect_client=request.args.get('client_id'),
                           today=date.today().isoformat())

@app.route('/taches/<int:id>/modifier', methods=['GET', 'POST'])
def tache_modifier(id):
    conn = get_db()
    tache = conn.execute("SELECT * FROM taches WHERE id=?", (id,)).fetchone()
    if request.method == 'POST':
        date_cloture = date.today().isoformat() if request.form.get('statut') == 'terminee' else None
        conn.execute('''UPDATE taches SET titre=?, description=?, client_id=?, contrat_id=?,
            date_echeance=?, priorite=?, statut=?, date_cloture=? WHERE id=?''',
            (request.form['titre'], request.form.get('description', ''),
             request.form.get('client_id') or None, request.form.get('contrat_id') or None,
             request.form.get('date_echeance', ''), request.form.get('priorite', 'normale'),
             request.form.get('statut', 'a_faire'), date_cloture, id))
        conn.commit()
        conn.close()
        flash('Tâche mise à jour !', 'success')
        return redirect(url_for('taches_list'))
    clients = conn.execute("SELECT id, nom, prenom FROM clients ORDER BY nom").fetchall()
    contrats = conn.execute("""SELECT ct.id, ct.numero, ct.type_assurance, c.nom
                               FROM contrats ct JOIN clients c ON ct.client_id=c.id
                               ORDER BY c.nom""").fetchall()
    conn.close()
    return render_template('taches/form.html', tache=tache, clients=clients, contrats=contrats,
                           titre='Modifier la tâche', preselect_client=None,
                           today=date.today().isoformat())

@app.route('/taches/<int:id>/terminer', methods=['POST'])
def tache_terminer(id):
    conn = get_db()
    conn.execute("UPDATE taches SET statut='terminee', date_cloture=? WHERE id=?",
                 (date.today().isoformat(), id))
    conn.commit()
    conn.close()
    flash('Tâche terminée ✓', 'success')
    return redirect(request.referrer or url_for('taches_list'))

@app.route('/taches/<int:id>/supprimer', methods=['POST'])
def tache_supprimer(id):
    conn = get_db()
    conn.execute("DELETE FROM taches WHERE id=?", (id,))
    conn.commit()
    conn.close()
    flash('Tâche supprimée.', 'info')
    return redirect(url_for('taches_list'))

# ─── API JSON ─────────────────────────────────────────────────────────────────

@app.route('/api/bateaux_client/<int:client_id>')
def api_bateaux_client(client_id):
    conn = get_db()
    bateaux = conn.execute(
        "SELECT id, nom_bateau, marque, modele FROM bateaux WHERE client_id=?",
        (client_id,)).fetchall()
    conn.close()
    return jsonify([dict(b) for b in bateaux])

@app.route('/api/contrats_client/<int:client_id>')
def api_contrats_client(client_id):
    conn = get_db()
    contrats = conn.execute(
        "SELECT id, numero, type_assurance FROM contrats WHERE client_id=?",
        (client_id,)).fetchall()
    conn.close()
    return jsonify([dict(c) for c in contrats])

# ─── MAIN ─────────────────────────────────────────────────────────────────────


# ═══════════════════════════════════════════════════
#  MODULES V4 — Paramètres, Alertes, Emails, Export,
#               Mandats, RGPD, DDA
# ═══════════════════════════════════════════════════

import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from flask import Response

def style_header(ws, row, cols, color="1a2e5a"):
    fill = PatternFill("solid", fgColor=color)
    font = Font(bold=True, color="FFFFFF", size=10)
    for col in range(1, cols+1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill; cell.font = font
        cell.alignment = Alignment(horizontal='center', vertical='center')

def style_alt(ws, row, cols, even=True):
    fill = PatternFill("solid", fgColor="EFF6FF" if even else "FFFFFF")
    for col in range(1, cols+1):
        ws.cell(row=row, column=col).fill = fill

# ─── PARAMÈTRES ──────────────────────────────────────────────

@app.route('/parametres', methods=['GET', 'POST'])
def parametres():
    if request.method == 'POST':
        conn = get_db()
        for key in request.form:
            conn.execute("INSERT OR REPLACE INTO parametres (cle,valeur) VALUES (?,?)", (key, request.form[key]))
        conn.commit(); conn.close()
        flash('Paramètres enregistrés !', 'success')
        return redirect(url_for('parametres'))
    params = get_all_params()
    return render_template('parametres/index.html', params=params)

# ─── ALERTES ÉCHÉANCES ────────────────────────────────────────

@app.route('/alertes')
def alertes():
    conn = get_db()
    today = date.today().isoformat()
    alertes_data = conn.execute("""
        SELECT ct.*, c.nom, c.prenom, c.email, c.telephone, b.nom_bateau,
               CAST(julianday(ct.date_fin) - julianday(?) AS INTEGER) as jours_restants
        FROM contrats ct JOIN clients c ON ct.client_id=c.id
        LEFT JOIN bateaux b ON ct.bateau_id=b.id
        WHERE ct.statut='en_cours' AND ct.date_fin IS NOT NULL AND ct.date_fin != ''
          AND julianday(ct.date_fin) - julianday(?) <= 90
          AND julianday(ct.date_fin) >= julianday(?)
        ORDER BY ct.date_fin ASC
    """, (today, today, today)).fetchall()
    expires = conn.execute("""
        SELECT ct.*, c.nom, c.prenom, c.email, b.nom_bateau
        FROM contrats ct JOIN clients c ON ct.client_id=c.id
        LEFT JOIN bateaux b ON ct.bateau_id=b.id
        WHERE ct.statut='en_cours' AND ct.date_fin IS NOT NULL AND ct.date_fin != ''
          AND julianday(ct.date_fin) < julianday(?) ORDER BY ct.date_fin DESC LIMIT 20
    """, (today,)).fetchall()
    conn.close()
    j30 = [a for a in alertes_data if a['jours_restants'] is not None and a['jours_restants'] <= 30]
    j60 = [a for a in alertes_data if a['jours_restants'] is not None and 30 < a['jours_restants'] <= 60]
    j90 = [a for a in alertes_data if a['jours_restants'] is not None and 60 < a['jours_restants'] <= 90]
    return render_template('alertes/index.html', j30=j30, j60=j60, j90=j90, expires=expires, today=today)

# ─── MODÈLES EMAILS ──────────────────────────────────────────

@app.route('/emails')
def emails_list():
    conn = get_db()
    templates = conn.execute("SELECT * FROM email_templates ORDER BY categorie, titre").fetchall()
    conn.close()
    cats = {}
    for t in templates:
        cats.setdefault(t['categorie'], []).append(t)
    return render_template('emails/index.html', categories=cats)

@app.route('/emails/nouveau', methods=['GET', 'POST'])
def email_nouveau():
    if request.method == 'POST':
        conn = get_db()
        conn.execute("INSERT INTO email_templates (categorie,titre,sujet,corps) VALUES (?,?,?,?)",
                     (request.form.get('categorie','Autre'), request.form['titre'],
                      request.form.get('sujet',''), request.form['corps']))
        conn.commit(); conn.close()
        flash('Modèle créé !', 'success')
        return redirect(url_for('emails_list'))
    return render_template('emails/form.html', template=None, titre='Nouveau modèle')

@app.route('/emails/<int:id>/modifier', methods=['GET', 'POST'])
def email_modifier(id):
    conn = get_db()
    tmpl = conn.execute("SELECT * FROM email_templates WHERE id=?", (id,)).fetchone()
    if request.method == 'POST':
        conn.execute("UPDATE email_templates SET categorie=?,titre=?,sujet=?,corps=? WHERE id=?",
                     (request.form.get('categorie','Autre'), request.form['titre'],
                      request.form.get('sujet',''), request.form['corps'], id))
        conn.commit(); conn.close()
        flash('Modèle mis à jour !', 'success')
        return redirect(url_for('emails_list'))
    conn.close()
    return render_template('emails/form.html', template=tmpl, titre='Modifier le modèle')

@app.route('/emails/<int:id>/supprimer', methods=['POST'])
def email_supprimer(id):
    conn = get_db()
    conn.execute("DELETE FROM email_templates WHERE id=?", (id,))
    conn.commit(); conn.close()
    flash('Modèle supprimé.', 'info')
    return redirect(url_for('emails_list'))



# ─── EXPORTS EXCEL ────────────────────────────────────────────

@app.route('/export/clients')
def export_clients():
    conn = get_db()
    rows = conn.execute("SELECT * FROM clients ORDER BY nom").fetchall()
    conn.close()
    wb = Workbook(); ws = wb.active; ws.title = "Clients"
    hdrs = ['ID','Type','Nom','Prénom','Email','Téléphone','Adresse','CP','Ville','SIREN','Statut','Créé le']
    ws.append(hdrs); style_header(ws, 1, len(hdrs))
    for i,r in enumerate(rows,2):
        ws.append([r['id'],r['type'],r['nom'],r['prenom'] or '',r['email'] or '',
                   r['telephone'] or '',r['adresse'] or '',r['code_postal'] or '',
                   r['ville'] or '',r['siren'] or '',r['statut'],r['date_creation'] or ''])
        style_alt(ws,i,len(hdrs),i%2==0)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 16
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return Response(buf.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename=RMA_clients_{date.today()}.xlsx'})

@app.route('/export/contrats')
def export_contrats():
    conn = get_db()
    rows = conn.execute("""SELECT ct.*,c.nom,c.prenom,b.nom_bateau FROM contrats ct
        JOIN clients c ON ct.client_id=c.id LEFT JOIN bateaux b ON ct.bateau_id=b.id
        ORDER BY ct.date_creation DESC""").fetchall()
    conn.close()
    wb = Workbook(); ws = wb.active; ws.title = "Contrats"
    hdrs = ['ID','N° Police','Client','Bateau','Produit','Compagnie','Début','Fin','Prime/an','Franchise','Statut']
    ws.append(hdrs); style_header(ws, 1, len(hdrs))
    for i,r in enumerate(rows,2):
        client = f"{r['prenom'] or ''} {r['nom']}".strip()
        ws.append([r['id'],r['numero'] or '',client,r['nom_bateau'] or '',
                   r['type_assurance'] or '',r['compagnie'] or '',
                   r['date_debut'] or '',r['date_fin'] or '',
                   r['prime_annuelle'],r['franchise'],r['statut']])
        style_alt(ws,i,len(hdrs),i%2==0)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 16
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return Response(buf.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename=RMA_contrats_{date.today()}.xlsx'})

@app.route('/export/commissions')
def export_commissions():
    conn = get_db()
    rows = conn.execute("""SELECT cm.*,ct.numero,ct.type_assurance,c.nom,c.prenom
        FROM commissions cm LEFT JOIN contrats ct ON cm.contrat_id=ct.id
        LEFT JOIN clients c ON ct.client_id=c.id ORDER BY cm.annee DESC,cm.mois DESC""").fetchall()
    conn.close()
    mois_n = ['Jan','Fév','Mar','Avr','Mai','Jun','Jul','Aoû','Sep','Oct','Nov','Déc']
    wb = Workbook(); ws = wb.active; ws.title = "Commissions"
    hdrs = ['ID','Compagnie','Contrat','Produit','Client','Mois','Année','Brut €','Taux %','Net €','Statut','Paiement']
    ws.append(hdrs); style_header(ws,1,len(hdrs))
    for i,r in enumerate(rows,2):
        client = f"{r['prenom'] or ''} {r['nom'] or ''}".strip()
        mois = mois_n[r['mois']-1] if r['mois'] else ''
        ws.append([r['id'],r['compagnie'] or '',r['numero'] or '',r['type_assurance'] or '',
                   client,mois,r['annee'],r['montant_brut'],r['taux'],r['montant_net'],
                   r['statut'],r['date_paiement'] or ''])
        style_alt(ws,i,len(hdrs),i%2==0)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 15
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return Response(buf.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename=RMA_commissions_{date.today()}.xlsx'})

@app.route('/export/bateaux')
def export_bateaux():
    conn = get_db()
    rows = conn.execute("SELECT b.*,c.nom,c.prenom FROM bateaux b JOIN clients c ON b.client_id=c.id ORDER BY c.nom").fetchall()
    conn.close()
    wb = Workbook(); ws = wb.active; ws.title = "Bateaux"
    hdrs = ['ID','Propriétaire','Nom bateau','Type','Marque','Modèle','Immat.','Année','Long.(m)','Valeur €','Port','Zone','Moteur','CV','Carburant']
    ws.append(hdrs); style_header(ws,1,len(hdrs))
    for i,r in enumerate(rows,2):
        client = f"{r['prenom'] or ''} {r['nom']}".strip()
        ws.append([r['id'],client,r['nom_bateau'] or '',r['type_bateau'] or '',
                   r['marque'] or '',r['modele'] or '',r['immatriculation'] or '',
                   r['annee_fabrication'] or '',r['longueur'] or '',r['valeur_assurance'],
                   r['port_attache'] or '',r['zone_navigation'] or '',
                   r['moteur_marque'] or '',r['moteur_puissance'] or '',r['moteur_carburant'] or ''])
        style_alt(ws,i,len(hdrs),i%2==0)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 15
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return Response(buf.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename=RMA_bateaux_{date.today()}.xlsx'})

# ─── MANDATS ─────────────────────────────────────────────────

@app.route('/mandats')
def mandats_list():
    conn = get_db()
    statut = request.args.get('statut','')
    q = "SELECT * FROM mandats WHERE 1=1"
    p = []
    if statut: q += " AND statut=?"; p.append(statut)
    q += " ORDER BY compagnie"
    mandats = conn.execute(q, p).fetchall()
    conn.close()
    return render_template('mandats/index.html', mandats=mandats, statut=statut)

@app.route('/mandats/nouveau', methods=['GET','POST'])
def mandat_nouveau():
    if request.method == 'POST':
        conn = get_db()
        conn.execute("""INSERT INTO mandats (compagnie,type_mandat,numero_mandat,contact_nom,
            contact_email,contact_telephone,date_debut,date_fin,produits,statut,notes)
            VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
            (get_compagnie_from_form(),request.form.get('type_mandat',''),
             request.form.get('numero_mandat',''),request.form.get('contact_nom',''),
             request.form.get('contact_email',''),request.form.get('contact_telephone',''),
             request.form.get('date_debut',''),request.form.get('date_fin',''),
             request.form.get('produits',''),request.form.get('statut','actif'),
             request.form.get('notes','')))
        conn.commit(); conn.close()
        flash('Mandat enregistré !', 'success')
        return redirect(url_for('mandats_list'))
    return render_template('mandats/form.html', mandat=None, titre='Nouveau mandat')

@app.route('/mandats/<int:id>/modifier', methods=['GET','POST'])
def mandat_modifier(id):
    conn = get_db()
    mandat = conn.execute("SELECT * FROM mandats WHERE id=?", (id,)).fetchone()
    if request.method == 'POST':
        conn.execute("""UPDATE mandats SET compagnie=?,type_mandat=?,numero_mandat=?,
            contact_nom=?,contact_email=?,contact_telephone=?,date_debut=?,date_fin=?,
            produits=?,statut=?,notes=? WHERE id=?""",
            (get_compagnie_from_form(),request.form.get('type_mandat',''),
             request.form.get('numero_mandat',''),request.form.get('contact_nom',''),
             request.form.get('contact_email',''),request.form.get('contact_telephone',''),
             request.form.get('date_debut',''),request.form.get('date_fin',''),
             request.form.get('produits',''),request.form.get('statut','actif'),
             request.form.get('notes',''),id))
        conn.commit(); conn.close()
        flash('Mandat mis à jour !', 'success')
        return redirect(url_for('mandats_list'))
    conn.close()
    return render_template('mandats/form.html', mandat=mandat, titre='Modifier le mandat')

@app.route('/mandats/<int:id>/supprimer', methods=['POST'])
def mandat_supprimer(id):
    conn = get_db()
    conn.execute("DELETE FROM mandats WHERE id=?", (id,))
    conn.commit(); conn.close()
    flash('Mandat supprimé.', 'info')
    return redirect(url_for('mandats_list'))

# ─── RGPD ────────────────────────────────────────────────────

@app.route('/rgpd')
def rgpd_list():
    conn = get_db()
    demandes = conn.execute("""SELECT d.*,c.nom,c.prenom,c.email FROM rgpd_demandes d
        JOIN clients c ON d.client_id=c.id ORDER BY d.date_demande DESC""").fetchall()
    stats = {
        'total_clients': conn.execute("SELECT COUNT(*) FROM clients").fetchone()[0],
        'consentements': conn.execute("SELECT COUNT(DISTINCT client_id) FROM rgpd_consentements WHERE consentement=1").fetchone()[0],
        'demandes_en_cours': conn.execute("SELECT COUNT(*) FROM rgpd_demandes WHERE statut='en_cours'").fetchone()[0],
        'demandes_traitees': conn.execute("SELECT COUNT(*) FROM rgpd_demandes WHERE statut='traite'").fetchone()[0],
    }
    conn.close()
    return render_template('rgpd/index.html', demandes=demandes, stats=stats)

@app.route('/rgpd/demande/nouvelle', methods=['GET','POST'])
def rgpd_demande_nouvelle():
    conn = get_db()
    if request.method == 'POST':
        conn.execute("""INSERT INTO rgpd_demandes (client_id,type_demande,date_demande,statut,notes)
            VALUES (?,?,?,?,?)""",
            (request.form['client_id'],request.form['type_demande'],
             request.form.get('date_demande',date.today().isoformat()),
             'en_cours',request.form.get('notes','')))
        conn.commit(); conn.close()
        flash('Demande RGPD enregistrée !', 'success')
        return redirect(url_for('rgpd_list'))
    clients = conn.execute("SELECT id,nom,prenom,email FROM clients ORDER BY nom").fetchall()
    conn.close()
    return render_template('rgpd/form_demande.html', clients=clients, today=date.today().isoformat())

@app.route('/rgpd/demande/<int:id>/traiter', methods=['POST'])
def rgpd_demande_traiter(id):
    conn = get_db()
    conn.execute("UPDATE rgpd_demandes SET statut='traite',date_traitement=? WHERE id=?",
                 (date.today().isoformat(), id))
    conn.commit(); conn.close()
    flash('Demande traitée.', 'success')
    return redirect(url_for('rgpd_list'))

@app.route('/rgpd/client/<int:client_id>')
def rgpd_client(client_id):
    conn = get_db()
    client = conn.execute("SELECT * FROM clients WHERE id=?", (client_id,)).fetchone()
    consentements = conn.execute("SELECT * FROM rgpd_consentements WHERE client_id=? ORDER BY date_consentement DESC", (client_id,)).fetchall()
    demandes = conn.execute("SELECT * FROM rgpd_demandes WHERE client_id=? ORDER BY date_demande DESC", (client_id,)).fetchall()
    conn.close()
    return render_template('rgpd/client.html', client=client, consentements=consentements, demandes=demandes, date=date)

@app.route('/rgpd/consentement/ajouter', methods=['POST'])
def rgpd_consentement_ajouter():
    conn = get_db()
    conn.execute("""INSERT INTO rgpd_consentements (client_id,type_traitement,consentement,date_consentement,source,notes)
        VALUES (?,?,?,?,?,?)""",
        (request.form['client_id'],request.form['type_traitement'],
         1 if request.form.get('consentement')=='1' else 0,
         request.form.get('date_consentement',date.today().isoformat()),
         request.form.get('source',''),request.form.get('notes','')))
    conn.commit()
    client_id = request.form['client_id']
    conn.close()
    flash('Consentement enregistré !', 'success')
    return redirect(url_for('rgpd_client', client_id=client_id))

# ─── DDA ─────────────────────────────────────────────────────

@app.route('/dda')
def dda_list():
    conn = get_db()
    recueils = conn.execute("""SELECT d.*,c.nom,c.prenom FROM dda_recueils d
        JOIN clients c ON d.client_id=c.id ORDER BY d.date_recueil DESC""").fetchall()
    conn.close()
    return render_template('dda/index.html', recueils=recueils)

@app.route('/dda/nouveau', methods=['GET','POST'])
def dda_nouveau():
    conn = get_db()
    if request.method == 'POST':
        conn.execute("""INSERT INTO dda_recueils (client_id,date_recueil,conseiller,
            objectif_assurance,type_navigation,zone_navigation,experience_navigation,
            nb_personnes_bord,valeur_bateau,garanties_souhaitees,garanties_exclues,
            budget_annuel,situation_assurance,sinistres_anterieurs,besoins_specifiques,
            recommandation,accepte_signature,statut,notes,contrat_id,date_signature,version)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (request.form['client_id'],request.form.get('date_recueil',date.today().isoformat()),
             request.form.get('conseiller',''),request.form.get('objectif_assurance',''),
             request.form.get('type_navigation',''),request.form.get('zone_navigation',''),
             request.form.get('experience_navigation',''),request.form.get('nb_personnes_bord',''),
             request.form.get('valeur_bateau',''),request.form.get('garanties_souhaitees',''),
             request.form.get('garanties_exclues',''),request.form.get('budget_annuel',''),
             request.form.get('situation_assurance',''),request.form.get('sinistres_anterieurs',''),
             request.form.get('besoins_specifiques',''),request.form.get('recommandation',''),
             1 if request.form.get('accepte_signature') else 0,
             request.form.get('statut','brouillon'),request.form.get('notes',''),
             request.form.get('contrat_id') or None,
             request.form.get('date_signature') or None,
             int(request.form.get('version') or 1)))
        conn.commit(); conn.close()
        flash('Recueil de besoins enregistré !', 'success')
        return redirect(url_for('dda_list'))
    clients = conn.execute("SELECT id,nom,prenom FROM clients ORDER BY nom").fetchall()
    contrats = conn.execute("""SELECT ct.id,ct.numero,ct.type_assurance,c.nom,c.prenom
        FROM contrats ct JOIN clients c ON ct.client_id=c.id
        ORDER BY c.nom""").fetchall()
    client_id = request.args.get('client_id')
    conn.close()
    return render_template('dda/form.html', recueil=None, clients=clients, contrats=contrats,
                           titre='Nouveau recueil', preselect=client_id,
                           today=date.today().isoformat(),
                           conseiller=get_param('cabinet_nom','Riviera Marine Assurances'))

@app.route('/dda/<int:id>')
def dda_detail(id):
    conn = get_db()
    recueil = conn.execute("""SELECT d.*,c.nom,c.prenom,c.email,c.telephone FROM dda_recueils d
        JOIN clients c ON d.client_id=c.id WHERE d.id=?""", (id,)).fetchone()
    conn.close()
    if not recueil: return redirect(url_for('dda_list'))
    params = get_all_params()
    return render_template('dda/detail.html', recueil=recueil, params=params)

@app.route('/dda/<int:id>/modifier', methods=['GET','POST'])
def dda_modifier(id):
    conn = get_db()
    recueil = conn.execute("SELECT * FROM dda_recueils WHERE id=?", (id,)).fetchone()
    if request.method == 'POST':
        conn.execute("""UPDATE dda_recueils SET client_id=?,date_recueil=?,conseiller=?,
            objectif_assurance=?,type_navigation=?,zone_navigation=?,experience_navigation=?,
            nb_personnes_bord=?,valeur_bateau=?,garanties_souhaitees=?,garanties_exclues=?,
            budget_annuel=?,situation_assurance=?,sinistres_anterieurs=?,besoins_specifiques=?,
            recommandation=?,accepte_signature=?,statut=?,notes=?,
            contrat_id=?,date_signature=?,version=? WHERE id=?""",
            (request.form['client_id'],request.form.get('date_recueil',date.today().isoformat()),
             request.form.get('conseiller',''),request.form.get('objectif_assurance',''),
             request.form.get('type_navigation',''),request.form.get('zone_navigation',''),
             request.form.get('experience_navigation',''),request.form.get('nb_personnes_bord',''),
             request.form.get('valeur_bateau',''),request.form.get('garanties_souhaitees',''),
             request.form.get('garanties_exclues',''),request.form.get('budget_annuel',''),
             request.form.get('situation_assurance',''),request.form.get('sinistres_anterieurs',''),
             request.form.get('besoins_specifiques',''),request.form.get('recommandation',''),
             1 if request.form.get('accepte_signature') else 0,
             request.form.get('statut','brouillon'),request.form.get('notes',''),
             request.form.get('contrat_id') or None,
             request.form.get('date_signature') or None,
             int(request.form.get('version') or 1),id))
        conn.commit(); conn.close()
        flash('Recueil mis à jour !', 'success')
        return redirect(url_for('dda_detail', id=id))
    clients = conn.execute("SELECT id,nom,prenom FROM clients ORDER BY nom").fetchall()
    contrats = conn.execute("""SELECT ct.id,ct.numero,ct.type_assurance,c.nom,c.prenom
        FROM contrats ct JOIN clients c ON ct.client_id=c.id ORDER BY c.nom""").fetchall()
    conn.close()
    return render_template('dda/form.html', recueil=recueil, clients=clients, contrats=contrats,
                           titre='Modifier le recueil', preselect=None,
                           today=date.today().isoformat(),
                           conseiller=get_param('cabinet_nom','Riviera Marine Assurances'))

@app.route('/dda/<int:id>/supprimer', methods=['POST'])
def dda_supprimer(id):
    conn = get_db()
    conn.execute("DELETE FROM dda_recueils WHERE id=?", (id,))
    conn.commit(); conn.close()
    flash('Recueil supprimé.', 'info')
    return redirect(url_for('dda_list'))


# ═══════════════════════════════════════════════════════════════
#  MODULES V5 — Journal, Relances, Checklist, Primes, Agenda,
#               Recherche globale, Lien document-contrat
# ═══════════════════════════════════════════════════════════════

STATUTS_CONTRAT = [
    ('en_cours',       'En cours',                    'badge-en_cours'),
    ('devis_envoye',   'Devis envoyé',                'badge-devis'),
    ('attente_pieces', 'En attente de pièces',        'badge-attente'),
    ('transmis_cie',   'Transmis à la compagnie',     'badge-attendu'),
    ('accepte_cie',    'Accepté — attente police',    'badge-accepte_cie'),
    ('a_renouveler',   'À renouveler',                'badge-renouveler'),
    ('suspendu',       'Suspendu / Impayé',           'badge-suspendu'),
    ('renouvele',      'Renouvelé',                   'badge-inactif'),
    ('resilie_client', 'Résilié (client)',             'badge-resilie'),
    ('resilie_cie',    'Résilié (compagnie)',          'badge-resilie'),
    ('echu',           'Échu',                        'badge-echu'),
]

# Transitions autorisées (contrôle de cohérence)
TRANSITIONS_CONTRAT = {
    'devis_envoye':   ['attente_pieces', 'transmis_cie', 'en_cours', 'echu'],
    'attente_pieces': ['transmis_cie', 'en_cours', 'echu'],
    'transmis_cie':   ['accepte_cie', 'attente_pieces', 'echu'],
    'accepte_cie':    ['en_cours'],
    'en_cours':       ['suspendu', 'a_renouveler', 'transmis_cie',
                       'resilie_client', 'resilie_cie', 'echu', 'renouvele'],
    'suspendu':       ['en_cours', 'resilie_cie', 'echu'],
    'a_renouveler':   ['en_cours', 'renouvele', 'echu'],
}
app.jinja_env.globals['STATUTS_CONTRAT'] = STATUTS_CONTRAT

# ─── RECHERCHE GLOBALE ────────────────────────────────────────

@app.route('/recherche')
def recherche():
    q = request.args.get('q','').strip()
    if len(q) < 2:
        return render_template('recherche.html', q=q, results=None)
    conn = get_db()
    like = f'%{q}%'
    clients = conn.execute(
        "SELECT id, nom, prenom, email, telephone, statut FROM clients "
        "WHERE nom LIKE ? OR prenom LIKE ? OR email LIKE ? OR telephone LIKE ? LIMIT 8",
        (like,like,like,like)).fetchall()
    contrats = conn.execute(
        """SELECT ct.id, ct.numero, ct.type_assurance, ct.compagnie, ct.statut,
           c.nom, c.prenom FROM contrats ct JOIN clients c ON ct.client_id=c.id
           WHERE ct.numero LIKE ? OR ct.compagnie LIKE ? OR ct.type_assurance LIKE ?
           OR c.nom LIKE ? LIMIT 8""",
        (like,like,like,like)).fetchall()
    bateaux = conn.execute(
        """SELECT b.id, b.nom_bateau, b.marque, b.modele, b.immatriculation,
           c.id as client_id, c.nom, c.prenom FROM bateaux b
           JOIN clients c ON b.client_id=c.id
           WHERE b.nom_bateau LIKE ? OR b.immatriculation LIKE ? OR b.marque LIKE ? LIMIT 6""",
        (like,like,like)).fetchall()
    sinistres = conn.execute(
        """SELECT s.id, s.numero, s.type_sinistre, s.statut, c.nom, c.prenom
           FROM sinistres s JOIN clients c ON s.client_id=c.id
           WHERE s.numero LIKE ? OR s.type_sinistre LIKE ? OR c.nom LIKE ? LIMIT 6""",
        (like,like,like)).fetchall()
    conn.close()
    total = len(clients)+len(contrats)+len(bateaux)+len(sinistres)
    return render_template('recherche.html', q=q,
                           clients=clients, contrats=contrats,
                           bateaux=bateaux, sinistres=sinistres, total=total)

# ─── JOURNAL D'ACTIVITÉ ──────────────────────────────────────

@app.route('/journal')
def journal_list():
    conn = get_db()
    type_c = request.args.get('type','')
    search = request.args.get('q','')
    q = """SELECT j.*, c.nom, c.prenom, ct.numero as contrat_num
           FROM journal_activites j JOIN clients c ON j.client_id=c.id
           LEFT JOIN contrats ct ON j.contrat_id=ct.id WHERE 1=1"""
    p = []
    if type_c: q += " AND j.type_contact=?"; p.append(type_c)
    if search:
        q += " AND (c.nom LIKE ? OR j.objet LIKE ? OR j.contenu LIKE ?)"
        p += [f'%{search}%']*3
    q += " ORDER BY j.date_activite DESC, j.id DESC LIMIT 100"
    activites = conn.execute(q, p).fetchall()
    conn.close()
    return render_template('journal/index.html', activites=activites,
                           type_filtre=type_c, search=search)

@app.route('/journal/nouveau', methods=['GET','POST'])
def journal_nouveau():
    conn = get_db()
    if request.method == 'POST':
        conn.execute("""INSERT INTO journal_activites
            (client_id,contrat_id,date_activite,type_contact,sens,
             objet,contenu,resultat,suite_a_donner,auteur)
            VALUES (?,?,?,?,?,?,?,?,?,?)""",
            (request.form['client_id'], request.form.get('contrat_id') or None,
             request.form['date_activite'], request.form['type_contact'],
             request.form.get('sens','sortant'), request.form['objet'],
             request.form.get('contenu',''), request.form.get('resultat',''),
             request.form.get('suite_a_donner',''),
             get_param('cabinet_nom','RIVIERA MARINE ASSURANCES')))
        conn.commit(); conn.close()
        flash('Échange enregistré !', 'success')
        client_id = request.form['client_id']
        return redirect(url_for('client_detail', id=client_id) + '#journal')
    clients = conn.execute("SELECT id,nom,prenom FROM clients ORDER BY nom").fetchall()
    contrats = conn.execute("""SELECT ct.id,ct.numero,ct.type_assurance,c.nom
        FROM contrats ct JOIN clients c ON ct.client_id=c.id ORDER BY c.nom""").fetchall()
    client_id = request.args.get('client_id')
    conn.close()
    return render_template('journal/form.html', clients=clients, contrats=contrats,
                           preselect=client_id, today=date.today().isoformat())

@app.route('/journal/<int:id>/supprimer', methods=['POST'])
def journal_supprimer(id):
    conn = get_db()
    j = conn.execute("SELECT client_id FROM journal_activites WHERE id=?", (id,)).fetchone()
    client_id = j['client_id'] if j else None
    conn.execute("DELETE FROM journal_activites WHERE id=?", (id,))
    conn.commit(); conn.close()
    flash('Échange supprimé.', 'info')
    return redirect(url_for('client_detail', id=client_id) + '#journal' if client_id else url_for('journal_list'))

# ─── RELANCES ────────────────────────────────────────────────

@app.route('/relances')
def relances_list():
    conn = get_db()
    statut = request.args.get('statut','a_faire')
    q = """SELECT r.*, c.nom, c.prenom, c.email, c.telephone,
           ct.numero, ct.type_assurance, ct.date_fin
           FROM relances r JOIN clients c ON r.client_id=c.id
           JOIN contrats ct ON r.contrat_id=ct.id WHERE 1=1"""
    p = []
    if statut: q += " AND r.statut=?"; p.append(statut)
    q += " ORDER BY r.date_prevue ASC"
    relances = conn.execute(q, p).fetchall()
    counts = {
        'a_faire': conn.execute("SELECT COUNT(*) FROM relances WHERE statut='a_faire'").fetchone()[0],
        'effectuee': conn.execute("SELECT COUNT(*) FROM relances WHERE statut='effectuee'").fetchone()[0],
        'sans_suite': conn.execute("SELECT COUNT(*) FROM relances WHERE statut='sans_suite'").fetchone()[0],
        'en_retard': conn.execute("SELECT COUNT(*) FROM relances WHERE statut='a_faire' AND date_prevue < ?", (date.today().isoformat(),)).fetchone()[0],
    }
    conn.close()
    return render_template('relances/index.html', relances=relances,
                           statut=statut, counts=counts, today=date.today().isoformat())

@app.route('/relances/nouvelle', methods=['GET','POST'])
def relance_nouvelle():
    conn = get_db()
    if request.method == 'POST':
        conn.execute("""INSERT INTO relances
            (contrat_id,client_id,type_relance,date_prevue,moyen,statut,message)
            VALUES (?,?,?,?,?,?,?)""",
            (request.form['contrat_id'], request.form['client_id'],
             request.form.get('type_relance','echeance'),
             request.form.get('date_prevue',''),
             request.form.get('moyen','email'),
             request.form.get('statut','a_faire'),
             request.form.get('message','')))
        conn.commit(); conn.close()
        flash('Relance planifiée !', 'success')
        return redirect(url_for('relances_list'))
    clients = conn.execute("SELECT id,nom,prenom FROM clients ORDER BY nom").fetchall()
    contrats = conn.execute("""SELECT ct.id,ct.numero,ct.type_assurance,ct.date_fin,c.nom
        FROM contrats ct JOIN clients c ON ct.client_id=c.id
        WHERE ct.statut IN ('en_cours','a_renouveler','attente_pieces') ORDER BY c.nom""").fetchall()
    client_id = request.args.get('client_id')
    contrat_id = request.args.get('contrat_id')
    conn.close()
    return render_template('relances/form.html', relance=None, clients=clients,
                           contrats=contrats, titre='Nouvelle relance',
                           today=date.today().isoformat(),
                           preselect_client=client_id, preselect_contrat=contrat_id)

@app.route('/relances/<int:id>/effectuer', methods=['POST'])
def relance_effectuer(id):
    conn = get_db()
    conn.execute("""UPDATE relances SET statut='effectuee', date_effectuee=?,
        reponse=? WHERE id=?""",
        (date.today().isoformat(), request.form.get('reponse',''), id))
    conn.commit(); conn.close()
    flash('Relance marquée comme effectuée ✓', 'success')
    return redirect(request.referrer or url_for('relances_list'))

@app.route('/relances/<int:id>/sans_suite', methods=['POST'])
def relance_sans_suite(id):
    conn = get_db()
    conn.execute("UPDATE relances SET statut='sans_suite' WHERE id=?", (id,))
    conn.commit(); conn.close()
    flash('Relance classée sans suite.', 'info')
    return redirect(request.referrer or url_for('relances_list'))

@app.route('/relances/<int:id>/supprimer', methods=['POST'])
def relance_supprimer(id):
    conn = get_db()
    conn.execute("DELETE FROM relances WHERE id=?", (id,))
    conn.commit(); conn.close()
    flash('Relance supprimée.', 'info')
    return redirect(url_for('relances_list'))

# ─── CHECKLIST PIÈCES & SIGNATURES ────────────────────────────

PIECES_DEFAUT = [
    "CNI / Passeport", "RIB", "Permis mer côtier", "Permis hauturier",
    "Acte de francisation", "Titre de navigation", "Bulletin de souscription signé",
    "Recueil de besoins DDA signé", "Carte grise / CG bateau", "Rapport d'expertise",
]

# Modèles de checklist par type de risque/produit
CHECKLIST_MODELES = {
    'bateau': {
        'label': '⛵ Bateau de plaisance',
        'obligatoires': [
            "CNI / Passeport assuré",
            "Acte de francisation (ou titre de navigation étranger)",
            "Certificat de jauge / carnet de circulation",
            "Permis mer (côtier ou hauturier selon zone)",
            "Rapport d'expertise (si bateau > 7 ans ou valeur > 30 000 €)",
            "Bulletin de souscription signé",
            "RIB (pour SEPA compagnie)",
            "Recueil DDA signé",
        ],
        'optionnels': [
            "Photos du bateau (4 faces)",
            "Facture d'achat ou évaluation",
            "Titre de propriété",
            "Relevé de sinistres (N-5 ans)",
            "Justificatif d'expérience navigation",
        ]
    },
    'jet_ski': {
        'label': '🏄 Jet Ski',
        'obligatoires': [
            "CNI / Passeport assuré",
            "Carte grise du jet ski",
            "Permis mer côtier (obligatoire > 6 CV)",
            "Bulletin de souscription signé",
            "RIB (pour SEPA compagnie)",
            "Recueil DDA signé",
        ],
        'optionnels': [
            "Facture d'achat",
            "Photos du jet ski",
        ]
    },
    'peniche': {
        'label': '🚢 Péniche / Fluvial',
        'obligatoires': [
            "CNI / Passeport assuré",
            "Titre de navigation fluviale",
            "Certificat de visite technique (VNF)",
            "Permis de conduire bateau de plaisance (CPCB ou équivalent)",
            "Bulletin de souscription signé",
            "RIB (pour SEPA compagnie)",
            "Recueil DDA signé",
        ],
        'optionnels': [
            "Rapport d'expertise / évaluation",
            "Certificat de conformité équipements",
            "Photos extérieures et intérieures",
            "Acte de propriété / acte de vente",
        ]
    },
    'yacht': {
        'label': '⚓ Yacht / Grande plaisance',
        'obligatoires': [
            "CNI / Passeport capitaine",
            "Brevet de capitaine (200 UMS minimum recommandé)",
            "Acte de francisation / certificat de pavillon",
            "Rapport d'expertise récent (obligatoire > 3 ans)",
            "Liste de l'équipage signé si personnel à bord",
            "Contrats de travail équipage (si applicable)",
            "Bulletin de souscription signé",
            "RIB (pour SEPA compagnie)",
            "Recueil DDA signé",
            "Inventaire de sécurité",
        ],
        'optionnels': [
            "Photos du yacht (extérieur + carré + cabines)",
            "Livre de bord dernière saison",
            "Relevé de sinistres (N-5 ans)",
            "Certificat de classe (Lloyd's, Bureau Veritas...)",
            "Certificat de navigabilité",
            "Évaluation notariale si valeur > 500 000 €",
        ]
    },
    'local_pro': {
        'label': '🏢 Local professionnel',
        'obligatoires': [
            "Extrait Kbis ou justificatif d'immatriculation",
            "Bail commercial ou titre de propriété",
            "Plan des locaux avec surfaces",
            "Attestation valeur du contenu professionnel",
            "Bulletin de souscription signé",
            "RIB",
            "Recueil DDA signé",
        ],
        'optionnels': [
            "Rapport de prévention incendie",
            "Certificat d'alarme intrusion",
            "Bilan comptable (N-1) si CA > 500 000 €",
            "Liste du matériel assuré",
        ]
    },
    'pro_nautisme': {
        'label': '🔧 Professionnel du nautisme',
        'obligatoires': [
            "Extrait Kbis",
            "Bilan comptable (N-1)",
            "Liste des bateaux de la flotte avec valeurs",
            "Attestation valeur des locaux / contenu",
            "Descriptif de l'activité",
            "Bulletin de souscription signé (RC)",
            "Bulletin de souscription signé (Flotte/Locaux si séparé)",
            "RIB",
            "Recueil DDA signé",
        ],
        'optionnels': [
            "Relevé de sinistres professionnels (N-5 ans)",
            "Certificats de qualification des techniciens",
            "Plan des locaux et quais",
            "Contrats de concession portuaire",
        ]
    },
    'habitation': {
        'label': '🏠 Habitation',
        'obligatoires': [
            "CNI / Passeport assuré",
            "Titre de propriété ou bail locatif",
            "Plan ou superficie certifiée",
            "Bulletin de souscription signé",
            "RIB",
            "Recueil DDA signé",
        ],
        'optionnels': [
            "Rapport d'expertise immobilière",
            "Attestation alarme (réduction de prime)",
            "Photos du bien",
            "Inventaire objets de valeur",
            "Diagnostic technique (DPE, etc.)",
        ]
    },
    'default': {
        'label': '📋 Standard',
        'obligatoires': [
            "CNI / Passeport assuré",
            "Bulletin de souscription signé",
            "RIB",
            "Recueil DDA signé",
        ],
        'optionnels': [
            "Rapport d'expertise",
            "Relevé de sinistres",
        ]
    }
}

@app.route('/contrats/<int:id>/checklist', methods=['GET','POST'])
def contrat_checklist(id):
    conn = get_db()
    contrat = conn.execute("""SELECT ct.*,c.nom,c.prenom FROM contrats ct
        JOIN clients c ON ct.client_id=c.id WHERE ct.id=?""", (id,)).fetchone()
    if request.method == 'POST':
        action = request.form.get('action')
        if action == 'add':
            conn.execute("INSERT INTO checklist_pieces (contrat_id,libelle,obligatoire) VALUES (?,?,?)",
                         (id, request.form['libelle'], 1 if request.form.get('obligatoire') else 0))
        elif action == 'toggle':
            piece_id = request.form['piece_id']
            recu = 1 if request.form.get('recu') else 0
            date_r = date.today().isoformat() if recu else None
            conn.execute("UPDATE checklist_pieces SET recu=?,date_reception=? WHERE id=?",
                         (recu, date_r, piece_id))
        elif action == 'delete':
            conn.execute("DELETE FROM checklist_pieces WHERE id=? AND contrat_id=?",
                         (request.form['piece_id'], id))
        elif action == 'init':
            # Initialise selon le type de risque du contrat
            modele_key = request.form.get('modele', 'default')
            modele = CHECKLIST_MODELES.get(modele_key, CHECKLIST_MODELES['default'])
            for p in modele['obligatoires']:
                conn.execute("INSERT OR IGNORE INTO checklist_pieces (contrat_id,libelle,obligatoire) VALUES (?,?,1)",
                             (id, p))
            if request.form.get('avec_optionnels'):
                for p in modele.get('optionnels', []):
                    conn.execute("INSERT OR IGNORE INTO checklist_pieces (contrat_id,libelle,obligatoire) VALUES (?,?,0)",
                                 (id, p))
        # Update completeness flag
        total = conn.execute("SELECT COUNT(*) FROM checklist_pieces WHERE contrat_id=? AND obligatoire=1", (id,)).fetchone()[0]
        recus = conn.execute("SELECT COUNT(*) FROM checklist_pieces WHERE contrat_id=? AND obligatoire=1 AND recu=1", (id,)).fetchone()[0]
        done = 1 if total > 0 and total == recus else 0
        conn.execute("UPDATE contrats SET checklist_done=? WHERE id=?", (done, id))
        conn.commit()
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            conn.close()
            return jsonify({'ok': True, 'done': done})
        conn.close()
        return redirect(url_for('contrat_checklist', id=id))
    pieces = conn.execute("SELECT * FROM checklist_pieces WHERE contrat_id=? ORDER BY obligatoire DESC, id", (id,)).fetchall()
    # Detect risque type from linked risque or bateau
    risque_type = 'default'
    if contrat['risque_id']:
        r = conn.execute("SELECT type_risque FROM risques WHERE id=?", (contrat['risque_id'],)).fetchone()
        if r: risque_type = r['type_risque']
    elif contrat['bateau_id']:
        risque_type = 'bateau'
    conn.close()
    return render_template('contrats/checklist.html', contrat=contrat, pieces=pieces,
                           pieces_defaut=PIECES_DEFAUT,
                           modeles=CHECKLIST_MODELES, risque_type=risque_type)

# ─── PRIMES PAIEMENTS & ÉCHEANCIER ────────────────────────────

@app.route('/contrats/<int:id>/primes', methods=['GET','POST'])
def contrat_primes(id):
    conn = get_db()
    contrat = conn.execute("""SELECT ct.*,c.nom,c.prenom FROM contrats ct
        JOIN clients c ON ct.client_id=c.id WHERE ct.id=?""", (id,)).fetchone()
    if request.method == 'POST':
        action = request.form.get('action','add')
        if action == 'add':
            conn.execute("""INSERT INTO primes_paiements
                (contrat_id,fraction,montant,date_echeance,mode_paiement,statut,reference,notes)
                VALUES (?,?,?,?,?,?,?,?)""",
                (id, request.form.get('fraction',1),
                 float(request.form.get('montant',0)),
                 request.form.get('date_echeance',''),
                 request.form.get('mode_paiement','virement'),
                 request.form.get('statut','attendu'),
                 request.form.get('reference',''),
                 request.form.get('notes','')))
        elif action == 'payer':
            conn.execute("""UPDATE primes_paiements SET statut='paye',
                date_paiement=?,mode_paiement=? WHERE id=? AND contrat_id=?""",
                (date.today().isoformat(),
                 request.form.get('mode_paiement','virement'),
                 request.form['prime_id'], id))
        elif action == 'generer':
            # Génère l'écheancier automatiquement
            prime_an = contrat['prime_annuelle']
            nb = int(request.form.get('nb_fractions', 1))
            montant_f = round(prime_an / nb, 2)
            date_debut = contrat['date_debut'] or date.today().isoformat()
            import datetime
            d = datetime.date.fromisoformat(date_debut)
            # Supprimer les anciennes
            conn.execute("DELETE FROM primes_paiements WHERE contrat_id=?", (id,))
            for i in range(nb):
                if nb == 1: months = 0
                elif nb == 2: months = i * 6
                elif nb == 4: months = i * 3
                else: months = i
                from datetime import date as dt
                import calendar
                m = d.month + months
                y = d.year + (m-1)//12
                m = ((m-1)%12)+1
                day = min(d.day, calendar.monthrange(y,m)[1])
                ech = datetime.date(y,m,day).isoformat()
                conn.execute("""INSERT INTO primes_paiements
                    (contrat_id,fraction,montant,date_echeance,mode_paiement,statut)
                    VALUES (?,?,?,?,?,'attendu')""",
                    (id, i+1, montant_f, ech,
                     request.form.get('mode_paiement','virement')))
            conn.execute("UPDATE contrats SET nb_fractions=?,mode_paiement=? WHERE id=?",
                         (nb, request.form.get('mode_paiement','virement'), id))
        elif action == 'delete':
            conn.execute("DELETE FROM primes_paiements WHERE id=? AND contrat_id=?",
                         (request.form['prime_id'], id))
        conn.commit()
        conn.close()
        return redirect(url_for('contrat_primes', id=id))
    primes = conn.execute("""SELECT * FROM primes_paiements WHERE contrat_id=?
        ORDER BY fraction""", (id,)).fetchall()
    total_attendu = sum(p['montant'] for p in primes)
    total_paye = sum(p['montant'] for p in primes if p['statut']=='paye')
    conn.close()
    return render_template('contrats/primes.html', contrat=contrat, primes=primes,
                           total_attendu=total_attendu, total_paye=total_paye,
                           today=date.today().isoformat())

# ─── AGENDA / RENDEZ-VOUS ────────────────────────────────────

@app.route('/agenda')
def agenda():
    conn = get_db()
    mois = int(request.args.get('mois', date.today().month))
    annee = int(request.args.get('annee', date.today().year))
    import calendar
    cal = calendar.monthcalendar(annee, mois)
    rdvs = conn.execute("""SELECT r.*,c.nom,c.prenom FROM rendez_vous r
        LEFT JOIN clients c ON r.client_id=c.id
        WHERE strftime('%Y-%m', r.date_rdv) = ?
        ORDER BY r.date_rdv, r.heure_debut""",
        (f'{annee:04d}-{mois:02d}',)).fetchall()
    # Group by day
    rdv_par_jour = {}
    for r in rdvs:
        day = int(r['date_rdv'].split('-')[2])
        rdv_par_jour.setdefault(day, []).append(r)
    # Upcoming
    a_venir = conn.execute("""SELECT r.*,c.nom,c.prenom FROM rendez_vous r
        LEFT JOIN clients c ON r.client_id=c.id
        WHERE r.date_rdv >= ? AND r.statut != 'annule'
        ORDER BY r.date_rdv, r.heure_debut LIMIT 10""",
        (date.today().isoformat(),)).fetchall()
    conn.close()
    mois_noms = ['Janvier','Février','Mars','Avril','Mai','Juin',
                  'Juillet','Août','Septembre','Octobre','Novembre','Décembre']
    prev_m = mois-1 if mois > 1 else 12
    prev_y = annee if mois > 1 else annee-1
    next_m = mois+1 if mois < 12 else 1
    next_y = annee if mois < 12 else annee+1
    return render_template('agenda/index.html', cal=cal, rdv_par_jour=rdv_par_jour,
                           mois=mois, annee=annee, mois_nom=mois_noms[mois-1],
                           prev_mois=prev_m, prev_annee=prev_y,
                           next_mois=next_m, next_annee=next_y,
                           a_venir=a_venir, today=date.today().isoformat())

@app.route('/agenda/nouveau', methods=['GET','POST'])
def rdv_nouveau():
    conn = get_db()
    if request.method == 'POST':
        conn.execute("""INSERT INTO rendez_vous
            (titre,client_id,contrat_id,date_rdv,heure_debut,heure_fin,
             lieu,type_rdv,statut,notes)
            VALUES (?,?,?,?,?,?,?,?,?,?)""",
            (request.form['titre'],
             request.form.get('client_id') or None,
             request.form.get('contrat_id') or None,
             request.form['date_rdv'],
             request.form.get('heure_debut',''),
             request.form.get('heure_fin',''),
             request.form.get('lieu',''),
             request.form.get('type_rdv','appel'),
             request.form.get('statut','prevu'),
             request.form.get('notes','')))
        conn.commit(); conn.close()
        flash('Rendez-vous créé !', 'success')
        return redirect(url_for('agenda'))
    clients = conn.execute("SELECT id,nom,prenom FROM clients ORDER BY nom").fetchall()
    contrats = conn.execute("""SELECT ct.id,ct.numero,ct.type_assurance,c.nom
        FROM contrats ct JOIN clients c ON ct.client_id=c.id ORDER BY c.nom""").fetchall()
    rdv_date = request.args.get('date', date.today().isoformat())
    conn.close()
    return render_template('agenda/form.html', rdv=None, clients=clients,
                           contrats=contrats, titre='Nouveau rendez-vous',
                           rdv_date=rdv_date)

@app.route('/agenda/<int:id>/modifier', methods=['GET','POST'])
def rdv_modifier(id):
    conn = get_db()
    rdv = conn.execute("SELECT * FROM rendez_vous WHERE id=?", (id,)).fetchone()
    if request.method == 'POST':
        conn.execute("""UPDATE rendez_vous SET titre=?,client_id=?,contrat_id=?,
            date_rdv=?,heure_debut=?,heure_fin=?,lieu=?,type_rdv=?,statut=?,notes=?
            WHERE id=?""",
            (request.form['titre'],
             request.form.get('client_id') or None,
             request.form.get('contrat_id') or None,
             request.form['date_rdv'],
             request.form.get('heure_debut',''),
             request.form.get('heure_fin',''),
             request.form.get('lieu',''),
             request.form.get('type_rdv','appel'),
             request.form.get('statut','prevu'),
             request.form.get('notes',''), id))
        conn.commit(); conn.close()
        flash('RDV mis à jour !', 'success')
        return redirect(url_for('agenda'))
    clients = conn.execute("SELECT id,nom,prenom FROM clients ORDER BY nom").fetchall()
    contrats = conn.execute("""SELECT ct.id,ct.numero,ct.type_assurance,c.nom
        FROM contrats ct JOIN clients c ON ct.client_id=c.id ORDER BY c.nom""").fetchall()
    conn.close()
    return render_template('agenda/form.html', rdv=rdv, clients=clients,
                           contrats=contrats, titre='Modifier le rendez-vous',
                           rdv_date=rdv['date_rdv'] if rdv else date.today().isoformat())

@app.route('/agenda/<int:id>/supprimer', methods=['POST'])
def rdv_supprimer(id):
    conn = get_db()
    conn.execute("DELETE FROM rendez_vous WHERE id=?", (id,))
    conn.commit(); conn.close()
    flash('RDV supprimé.', 'info')
    return redirect(url_for('agenda'))

# ─── API: journal par client (AJAX) ──────────────────────────

@app.route('/api/journal_client/<int:client_id>')
def api_journal_client(client_id):
    conn = get_db()
    activites = conn.execute("""SELECT j.*,ct.numero FROM journal_activites j
        LEFT JOIN contrats ct ON j.contrat_id=ct.id
        WHERE j.client_id=? ORDER BY j.date_activite DESC, j.id DESC LIMIT 50""",
        (client_id,)).fetchall()
    conn.close()
    return jsonify([dict(a) for a in activites])

# ─── Enrichissement route contrat detail ─────────────────────

@app.route('/contrats/<int:id>')
def contrat_detail(id):
    conn = get_db()
    contrat = conn.execute("""SELECT ct.*,c.nom,c.prenom,c.email,c.telephone,
        b.nom_bateau,b.marque,b.immatriculation FROM contrats ct
        JOIN clients c ON ct.client_id=c.id
        LEFT JOIN bateaux b ON ct.bateau_id=b.id WHERE ct.id=?""", (id,)).fetchone()
    if contrat:
        contrat = dict(contrat)
    if not contrat:
        flash('Contrat introuvable.', 'error')
        return redirect(url_for('contrats_list'))
    pieces = conn.execute("SELECT * FROM checklist_pieces WHERE contrat_id=? ORDER BY obligatoire DESC,id", (id,)).fetchall()
    primes = conn.execute("SELECT * FROM primes_paiements WHERE contrat_id=? ORDER BY fraction", (id,)).fetchall()
    relances = conn.execute("""SELECT r.* FROM relances r WHERE r.contrat_id=?
        ORDER BY r.date_prevue DESC""", (id,)).fetchall()
    docs = conn.execute("SELECT * FROM documents WHERE contrat_id=? ORDER BY date_upload DESC", (id,)).fetchall()
    journal = conn.execute("""SELECT j.*,c.nom,c.prenom FROM journal_activites j
        JOIN clients c ON j.client_id=c.id WHERE j.contrat_id=?
        ORDER BY j.date_activite DESC LIMIT 10""", (id,)).fetchall()
    # NAVISUR v9.4 — Suivi financier
    primes_hist = conn.execute(
        "SELECT * FROM primes_historique WHERE contrat_id=? ORDER BY annee DESC", (id,)).fetchall()
    quittances_sf = conn.execute(
        """SELECT * FROM quittances_suivi WHERE contrat_id=?
           ORDER BY annee DESC, numero_quittance ASC""", (id,)).fetchall()
    cal_prelev = {}
    for q in quittances_sf:
        if q['mode_paiement'] == 'prelevement':
            cal_prelev[q['id']] = conn.execute(
                "SELECT * FROM calendrier_prelevements WHERE quittance_id=? ORDER BY numero_echeance",
                (q['id'],)).fetchall()
    conn.close()
    
    total_paye_quitt = sum(q['montant_ttc'] for q in quittances_sf if q['statut_paiement'] == 'payee')
    total_attendu_quitt = sum(q['montant_ttc'] for q in quittances_sf)
    total_paye = total_paye_quitt if total_paye_quitt > 0 else total_paye
    total_attendu = total_attendu_quitt if total_attendu_quitt > 0 else total_attendu
    # Stats suivi financier
    nb_quitt_non_envoyees = sum(1 for q in quittances_sf if not q['envoyee_client'] and q['mode_paiement'] != 'prelevement')
    nb_quitt_impayees = sum(1 for q in quittances_sf if q['statut_paiement'] == 'impayee')
    comm_attendue_total = sum(q['commission_attendue'] or 0 for q in quittances_sf)
    comm_recue_total = sum(q['commission_recue'] or 0 for q in quittances_sf)
    conn_audit = get_db()
    audit_logs = conn_audit.execute("SELECT * FROM audit_log WHERE table_name='contrats' AND record_id=? ORDER BY date_action DESC", (id,)).fetchall()
    conn_audit.close()
    breadcrumbs = [
        {'label': 'Contrats', 'url': url_for('contrats_list')},
        {'label': f"Client: {contrat.get('prenom', '')} {contrat.get('nom', '')}".strip(), 'url': url_for('client_detail', id=contrat['client_id'])},
        {'label': f"Contrat {contrat.get('numero', id)}", 'url': None}
    ]
    return render_template('contrats/detail.html', contrat=contrat,
                           pieces=pieces, primes=primes, relances=relances,
                           docs=docs, journal=journal, sinistres_ct=sinistres_ct,
                           total_paye=total_paye, total_attendu=total_attendu,
                           pieces_ok=pieces_ok, pieces_total=pieces_total,
                           primes_hist=primes_hist, quittances_sf=quittances_sf,
                           cal_prelev=cal_prelev,
                           nb_quitt_non_envoyees=nb_quitt_non_envoyees,
                           nb_quitt_impayees=nb_quitt_impayees,
                           comm_attendue_total=comm_attendue_total,
                           comm_recue_total=comm_recue_total,
                           STATUTS_CONTRAT=STATUTS_CONTRAT,
                           today=date.today().isoformat(), breadcrumbs=breadcrumbs,
                           audit_logs=audit_logs)


# ═══════════════════════════════════════════════════════════════
#  MODULE COMPTABILITÉ v6
#  Dashboard · Livre recettes · Commissions · URSSAF
# ═══════════════════════════════════════════════════════════════

from datetime import datetime as dt
import json as _json

NATURES_RECETTE = [
    ('commission_apport',     'Commission d\'apport (nouveau contrat)'),
    ('commission_renouvellement','Commission de renouvellement'),
    ('surcommission',          'Surcommission / Bonus volume'),
    ('frais_courtage',         'Frais de courtage'),
    ('honoraire_conseil',      'Honoraire de conseil'),
    ('autre',                  'Autre'),
]
MODES_REGLEMENT = ['virement','cheque','especes','prelevement','autre']
STATUTS_COMMISSION = [
    ('attendue',     '⏳ Attendue'),
    ('releve_recu',  '📄 Relevé reçu'),
    ('encaissee',    '✅ Encaissée'),
    ('retard',       '⚠️ En retard'),
    ('contestee',    '❌ Contestée'),
    ('annulee',      '🚫 Annulée'),
]

# ── helpers compta ──────────────────────────────────────────────

def get_ca_annee(annee):
    conn = get_db()
    r = conn.execute(
        "SELECT COALESCE(SUM(montant),0) FROM compta_recettes WHERE annee=?",
        (annee,)).fetchone()[0]
    conn.close()
    return float(r)

def get_ca_mois(annee, mois):
    conn = get_db()
    r = conn.execute(
        "SELECT COALESCE(SUM(montant),0) FROM compta_recettes WHERE annee=? AND mois=?",
        (annee, mois)).fetchone()[0]
    conn.close()
    return float(r)

def next_numero_ordre():
    conn = get_db()
    r = conn.execute("SELECT COALESCE(MAX(numero_ordre),0) FROM compta_recettes").fetchone()[0]
    conn.close()
    return int(r) + 1

# calc_urssaf → déplacé dans core.py
def urssaf_dates_limites(annee):
    """Retourne les dates limites URSSAF trimestrielles."""
    return {
        1: f'{annee}-04-30',
        2: f'{annee}-07-31',
        3: f'{annee}-10-31',
        4: f'{annee+1}-01-31',
    }

# ── DASHBOARD COMPTABLE ─────────────────────────────────────────

